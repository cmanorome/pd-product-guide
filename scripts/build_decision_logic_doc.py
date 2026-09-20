#!/usr/bin/env python3
"""Build a Word copy of RECOMMENDATION-DECISION-LOGIC.md for team testing."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x22, 0xB1, 0x4C)
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F3FBF5"
HEADER_BG = "22B14C"
FORMULA_BG = "F6F7F8"


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run(run, *, size=11, bold=False, color=INK, italic=False, name="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def add_para(doc, text, *, size=11, bold=False, color=INK, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color=GREEN)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.75)
        run = p.add_run(item)
        set_run(run, size=11)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.75)
        run = p.add_run(item)
        set_run(run, size=11)


def formula_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run("\n".join(lines))
    set_run(run, size=10, name="Consolas", color=INK)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), FORMULA_BG)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run(run, size=10, bold=True, color=WHITE)
        shade(cell, HEADER_BG)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = tbl.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run(run, size=10)
            if r_i % 2 == 1:
                shade(cell, ROW_ALT)
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


def build() -> Path:
    root = Path(__file__).resolve().parents[1]
    out = root / "Plant-Doctor-Recommendation-Decision-Logic.docx"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("Plant Doctor Product Guide")
    set_run(r, size=22, bold=True, color=GREEN)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(4)
    r = sub.add_run("How recommendations are decided")
    set_run(r, size=16, bold=True, color=INK)

    add_para(
        doc,
        "Team testing notes — what the engine is trying to do, not every code path.",
        size=11,
        color=MUTED,
        space_after=4,
    )
    add_para(
        doc,
        "The live plan is meant to stay short. Open “Why we chose these” at the bottom of a result for a plain-language recap of that run.",
        space_after=10,
    )

    heading(doc, "1. Two modes")
    table(
        doc,
        ["Mode in the form", "What it optimises for"],
        [
            [
                "Fix a problem",
                "Symptoms + soil + where it is used. Intent is always “fix what’s wrong”.",
            ],
            [
                "Improve results",
                "Lawn, garden, or farm goals. Symptoms are ignored. “Where are you using it?” picks which goal set you see.",
            ],
        ],
        col_widths=[4.5, 12.0],
    )
    add_para(
        doc,
        "“Where are you using it?” is the main switch between lawn and garden. Lawn stays the primary business path. Garden beds / pots / indoor unlock garden products instead of turf specialists.",
    )

    heading(doc, "2. How a result is built")
    numbered(
        doc,
        [
            "Read the form (symptoms or goals, soil, season, optional soil test).",
            "Score every product from the CSV catalog.",
            "Drop products that don’t fit (hard constraints — see section 6).",
            "Pick a primary (highest valid score).",
            "Add supporting layers only when that role is warranted (section 5).",
            "Show matching kits under “Or use a kit” — kits never replace the step-by-step plan.",
            "On lawns, sometimes show Champion Fairway vs Greens Grade as a choice, not a score split.",
        ],
    )
    add_para(
        doc,
        "Typical stack order (after the hero product): Soil structure → soil pH → biology (seaweed) → uptake (humic / Stimulizer) → feed → colour (iron).",
        space_before=6,
    )
    add_para(
        doc,
        "Default cap is 4 products (5 if nutrient lockout and a structure issue are both on). Biology (seaweed) and uptake are almost always allowed as support layers. Gypsum, lime, wetter, zeolite, fertiliser, and iron only appear when the situation calls for them.",
    )

    heading(doc, "3. Scoring weights")
    add_para(doc, "Each product gets a base score, then a role multiplier.")
    heading(doc, "Symptom mode (Fix a problem)", level=2)
    formula_block(
        doc,
        [
            "base =",
            "  1.40 × problem match",
            "+ 0.90 × soil match",
            "+ 0.85 × use-case match",
            "+ 0.15 × seasonal match",
            "+ 0.08 × synergy",
            "+ targeted bonuses",
            "+ lawn/garden context fit",
            "",
            "final = base × role multiplier",
        ],
    )
    add_para(
        doc,
        "Problem / soil / season numbers come from the CSV (usually 0–5). Use-case columns are 0 or 1 (Lawn, Garden Beds, Pots, Indoor, Farms).",
    )
    add_para(
        doc,
        "Problem match is the average of the CSV scores for the symptoms the user ticked. Tick yellowing only → use the Yellowing Leaves column. Tick yellowing + slow growth → average of those two columns.",
    )

    heading(doc, "Goals mode (Improve results)", level=2)
    add_para(doc, "Goals dominate. Context (soil, place, season, symptoms) is a light overlay:")
    formula_block(
        doc,
        [
            "base =",
            "  1.20 × goal CSV fit × goal-alignment multiplier",
            "+ 0.15 × 0.85 × use-case",
            "+ 0.12 × 0.15 × season",
            "+ 0.28 × 0.90 × soil",
            "+ 0.05 × 0.08 × synergy",
            "+ 0.05 × 1.40 × problem   (usually 0 — symptoms are off)",
            "+ pair synergy (small)",
            "+ garden stage bonus",
            "+ lawn/garden context fit",
            "",
            "final = base × role multiplier",
        ],
    )
    add_para(
        doc,
        "Goal-alignment multiplier maps CSV goal fit (0–5) onto about 0.94–1.28. A product that scores 5/5 on the selected goals gets a lift; a weak fit does not.",
    )
    add_para(
        doc,
        "If several goals are ticked, no single goal can take more than 42% of the effective weight.",
    )

    heading(doc, "4. What “success” means by vertical")
    add_para(
        doc,
        "These are the built-in priorities when more than one goal is ticked. A single ticked goal still drives ranking via that product’s CSV column.",
    )
    heading(doc, "Lawn", level=2)
    table(
        doc,
        ["Goal", "Share"],
        [
            ["Thickening and density", "28%"],
            ["Fast recovery from stress", "22%"],
            ["Deep green colour", "18%"],
            ["Weed suppression through dominance", "16%"],
            ["Low maintenance resilience", "16%"],
        ],
        col_widths=[12.0, 4.0],
    )
    add_para(
        doc,
        "Colour is real, but density and recovery outrank colour so iron-only answers don’t win every lawn goal.",
    )
    heading(doc, "Garden", level=2)
    table(
        doc,
        ["Goal", "Share"],
        [
            ["Improved soil fertility over time", "28%"],
            ["Root development and transplant success", "26%"],
            ["Strong flowering and fruiting", "18%"],
            ["Pest and disease resilience", "14%"],
            ["Consistent growth across seasons", "14%"],
        ],
        col_widths=[12.0, 4.0],
    )
    heading(doc, "Farm", level=2)
    table(
        doc,
        ["Goal", "Share"],
        [
            ["Soil efficiency", "28%"],
            ["Yield increase", "22%"],
            ["Water efficiency", "20%"],
            ["Crop uniformity", "16%"],
            ["Reduced input dependency", "14%"],
        ],
        col_widths=[12.0, 4.0],
    )
    heading(doc, "Stage (goals mode only)", level=2)
    add_para(doc, "Shown as Just planting / Keep it healthy / Push for best results.")
    bullets(
        doc,
        [
            "Planting: extra weight on roots, soil fertility, density, yield.",
            "Keep it healthy: slight lift on low-maintenance / consistent growth.",
            "Best results: extra weight on colour, density, flowering, yield, uniformity.",
        ],
    )
    add_para(
        doc,
        "Symptom mode does not use this dropdown; it always scores as “fix the problem”.",
        space_before=6,
    )
    add_para(
        doc,
        "Complementary pairs (small bonus if both goals are ticked and the product scores 3 or more on both): roots + soil fertility; water efficiency + soil efficiency; density + colour; weed suppression + density.",
    )

    heading(doc, "5. Role multipliers")
    add_para(
        doc,
        "Applied after the base score. This is why a “pretty good” iron product can beat a “pretty good” fertiliser when yellowing is ticked.",
    )
    table(
        doc,
        ["Role", "When", "Multiplier"],
        [
            ["Soil structure (gypsum, wetter, zeolite)", "Not maintenance", "1.30"],
            ["Soil structure", "Maintenance", "1.22"],
            ["Soil chemistry (lime / dolomite)", "Acidic, alkaline, or lockout", "1.28"],
            ["Soil chemistry", "Otherwise", "1.05"],
            ["Biology (seaweed)", "Always", "1.20"],
            ["Uptake (humic / fulvic / Stimulizer)", "Always", "1.15"],
            ["Nutrition (fertiliser)", "Goals mode", "1.12"],
            ["Nutrition", "Slow growth, no lockout", "1.25"],
            ["Nutrition", "Other symptom cases", "0.95"],
            ["Visual (iron colour)", "Yellowing ticked", "1.20"],
            ["Visual", "No yellowing", "0.70"],
            ["Kit / bundle", "Scoring only — not used as the plan", "1.10"],
        ],
        col_widths=[7.5, 6.5, 2.5],
    )
    add_para(
        doc,
        "High-nitrogen products (MaxGreen, Activ8EXTRA) are cut 15% in summer if fungal issues are ticked.",
    )
    heading(doc, "When a supporting role is allowed in the stack", level=2)
    table(
        doc,
        ["Role", "Added when"],
        [
            ["Soil structure", "Clay, sandy, hydrophobic, compaction, or poor water holding"],
            ["Soil chemistry", "Acidic, alkaline, or nutrient lockout"],
            ["Biology", "Always"],
            ["Uptake", "Always"],
            [
                "Nutrition (symptom mode)",
                "Slow growth, patchy lawn, weak roots, fungal issues, or poor flowering",
            ],
            ["Nutrition (goals mode)", "Shown as “Feed with”, not mixed into the same stack loop"],
            ["Visual / colour", "Yellowing"],
        ],
        col_widths=[5.5, 11.0],
    )

    heading(doc, "6. Hard rules (can score well and still be excluded)")
    add_para(doc, "These are pass/fail, not weights.", bold=True)
    add_para(doc, "Place", bold=True, space_before=6, space_after=4)
    bullets(
        doc,
        [
            "Flowers, Fruits & Roots is not a lawn feed.",
            "Turf specialists (Champion, Lawn Envy, MaxGreen, lawn kits) are not used for garden beds, pots, indoor, or farm nutrition.",
            "Granular bed/turf feeds (Roots, Shoots & Leaves; Champion; FFR granular) are not the default for pots / indoor.",
            "FFR in pots / indoor only if flowering is the job (flowering goal or “poor flowering” symptom).",
        ],
    )
    add_para(doc, "Soil chemistry", bold=True, space_before=8, space_after=4)
    bullets(
        doc,
        [
            "Lime / dolomite only if soil is acidic (checkbox or measured pH). Never on alkaline.",
            "Gypsum only for clay or compaction — not for “weak roots” alone.",
            "Wetter only for hydrophobic soil or poor water holding.",
            "Zeolite only for sandy, low organic matter, or lockout.",
        ],
    )
    add_para(doc, "Program safety", bold=True, space_before=8, space_after=4)
    bullets(
        doc,
        [
            "Do not stack two iron products.",
            "Do not mix lime with iron in the same plan.",
            "Nutrient lockout: do not lead with fertiliser; unlock with biology / uptake / chemistry first.",
            "Iron can be the hero when yellowing is the main issue and lockout is not.",
            "Liquid iron + seaweed / humic / wetter can share a program but not a tank. The UI should warn: apply iron on a different day. Stimulizer is the uptake that can tank-mix with iron.",
        ],
    )
    add_para(doc, "Biology pick", bold=True, space_before=8, space_after=4)
    bullets(
        doc,
        [
            "Default biology is Seaweed Secrets, not neem.",
            "Neem is used when fungal issues or the garden pest/disease goal is on.",
        ],
    )

    heading(doc, "7. Extra bonuses (symptom mode)")
    add_para(doc, "On top of CSV scores, the engine nudges a few named jobs:")
    table(
        doc,
        ["Situation", "Nudge"],
        [
            ["Lawn + yellowing + iron product", "+0.90 (+1.15 if alkaline)"],
            ["Garden / general yellowing + iron", "+0.55"],
            ["Alkaline + yellowing + Kendon chelate", "extra +0.25"],
            ["Garden + poor flowering + FFR", "+1.05"],
            ["Garden + slow growth + Roots, Shoots & Leaves or Activ8Mate", "+0.40"],
            ["Garden + weak roots + Roots, Shoots & Leaves", "+0.90"],
            [
                "Garden + weak roots (no flowering) + FFR",
                "−0.45 (don’t use flower/fruit feed for transplant shock)",
            ],
        ],
        col_widths=[10.5, 6.0],
    )
    add_para(doc, "Lawn vs garden context fit", bold=True, space_before=6, space_after=4)
    bullets(
        doc,
        [
            "Garden place: garden-only lines (FFR, RSL, garden kits, most humics) +0.55. Activ8Mate +0.35. Activ8EXTRA −0.12. Turf specialists −1.0 (and they are already blocked).",
            "Lawn place: products with Lawn = 0 and Garden Beds = 1 get −0.55 (stops Roots, Shoots & Leaves stealing turf density goals).",
        ],
    )

    heading(doc, "8. Optional soil test")
    add_para(doc, "Leave blank if unused. A measured pH overrides acidic/alkaline checkboxes.")
    table(
        doc,
        ["Method", "Acidic", "Near-neutral", "Alkaline"],
        [
            ["Water (home kit)", "below ~6.0", "6.0–7.2", "above ~7.2"],
            ["CaCl₂ (AU lab)", "below ~5.3", "5.3–6.5", "above ~6.5"],
        ],
        col_widths=[4.5, 4.0, 4.0, 4.0],
    )
    add_para(
        doc,
        "Organic matter below 2% is treated as low OM (biology / carbon products become more relevant). A soil-test entry also slightly increases confidence.",
    )

    heading(doc, "9. Champion pair (lawns only)")
    add_para(
        doc,
        "Champion Fairway and Greens Grade are two intensities of the same turf fertiliser, not competing SKUs. The UI may show both:",
    )
    bullets(
        doc,
        [
            "Everyday lawns — Fairway",
            "Fine / low-cut — Greens Grade",
        ],
    )
    add_para(
        doc,
        "This block is off for garden beds, pots, and indoor. It appears when the user is in a lawn context and Champion is competitive vs the rest of the catalog (not on every lawn click). In lawn goals mode it must also be near the best goal-match in the catalog.",
        space_before=6,
    )

    heading(doc, "10. Kits")
    add_para(
        doc,
        "Kits are scored like other products but never become the default plan, even if many symptoms are ticked. They only appear under “Or use a kit”. Lawn kits for lawns; garden kits (Gardener’s Choice, soil enhancer packs) for garden users.",
    )

    heading(doc, "11. What testers should usually see")
    add_para(
        doc,
        "These are expected leads, not the only legal stack. Supporting seaweed / humic / Stimulizer is normal.",
    )
    table(
        doc,
        ["Test", "Expect to start with (or feed with)"],
        [
            ["Lawn + yellowing", "Liquid Iron. Champion pair may also show. Kit fold, not the hero."],
            [
                "Lawn + yellowing + alkaline pH (~7.8 water)",
                "Kendon Iron Chelate preferred over sulphate. No lime.",
            ],
            [
                "Lawn + acidic pH (~5.2 water)",
                "Lime/dolomite allowed if chemistry is in play; skip extra lime if pH is already OK.",
            ],
            [
                "Lawn + thickening / density goal",
                "Champion pair as the fertiliser choice; not Flowers, Fruits & Roots.",
            ],
            [
                "Garden beds + yellowing",
                "Liquid iron is still valid (chlorosis). Kits should be garden kits, not Lawn Lovers. No Champion.",
            ],
            ["Garden beds + slow growth", "Roots, Shoots & Leaves (or Activ8Mate)."],
            [
                "Garden beds + poor flowering, or flowering goal",
                "Flowers, Fruits & Roots liquid.",
            ],
            ["Garden beds + planting + root goal", "Roots, Shoots & Leaves leads."],
            ["Garden beds + richer soil goal", "Humic / seaweed first; Activ8Mate as the feed."],
            ["Pots + flowering", "FFR liquid, not granular RSL/FFR."],
            ["Clay + compaction (any place)", "Gypsum is allowed."],
            [
                "Weak roots without clay/compaction",
                "Not gypsum; garden → seaweed / RSL style feed.",
            ],
            ["Nutrient lockout", "Don’t lead with NPK; biology/uptake/chemistry first."],
            ["Iron + seaweed in the same plan", "Warning to apply iron on a different day."],
        ],
        col_widths=[6.5, 10.0],
    )

    heading(doc, "12. How to review a run")
    numbered(
        doc,
        [
            "Check “Where are you using it?” — lawn vs garden changes the catalog more than any weighting.",
            "Read the hero + numbered steps. Roles should match the table in section 5.",
            "Open “Why we chose these” — it should mention the place, the symptoms or goals, and each step’s job.",
            "Open “Or use a kit” only if you are checking bundles.",
            "If a product feels wrong, ask: was it scored high, or was a better product blocked (section 6)? Blocked products cannot appear no matter how well they score.",
        ],
    )
    add_para(
        doc,
        "CSV columns testers care about: Yellowing Leaves, Slow Growth, Weak Roots, Nutrient Lockout, Patchy Lawn, Fungal Issues, Compacted Soil, Poor Water Retention, Lawn / Garden Beds / Pots / Indoor / Farms, the soil columns, season columns, and the goal columns for that vertical.",
        space_before=8,
    )
    add_para(
        doc,
        "Catalog file: plant_doctor_recommendation_engine_template.csv. Engine rules: pd_engine/. Changing a CSV cell changes ranking; changing code changes when a product is allowed in at all.",
        color=MUTED,
        space_before=4,
    )

    doc.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(path)
